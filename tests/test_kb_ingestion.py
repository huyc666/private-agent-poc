"""PLAN 2.1 文档入库管道测试（spec：docs/specs/2.1-kb-ingestion.md，AC1~AC8）。

- AC1 五类新格式（docx/pdf/xlsx/pptx/html）解析入库——测试内程序化生成最小样例
- AC2 .md/.txt 索引与检索行为不回退
- AC3 向量通道：注入确定性假向量验证落库与点积排序（离线确定性）
- AC4 RRF 融合符合公式预期 + 单通道退化
- AC5 语义问法命中（RUN_EMBED_TESTS 门控：真 fastembed 模型，需模型文件）
- AC6 拒绝路径：老格式报错 / 损坏文件 / 扫描件 / 超限 → 跳过 + WARNING 不阻塞
- AC7 全链路零网络：封禁 socket 下索引与查询跑通
- AC8 包检索工具自包含（源码静态断言，不 import server/tools）
"""
import array
import importlib.util
import json
import os
import socket
import sqlite3
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent

import tools.build_kb as build_kb  # noqa: E402
import tools.kb_parsers as kb_parsers  # noqa: E402

PACK_TOOLS = {
    "common": PROJECT_ROOT / "packs/core-knowledge/tools/search_common_knowledge.py",
    "hr": PROJECT_ROOT / "packs/kb-demo/tools/search_hr_knowledge.py",
}

# ---------------------------------------------------------------- 工具函数


def _build_kb(tmp_path, files, embed_fn="unset"):
    """落盘文档并建索引；embed_fn 用 sentinel 区分「未指定（走 env）」与「显式关闭」。"""
    kb = tmp_path / "kb"
    kb.mkdir(parents=True, exist_ok=True)
    for name, content in files.items():
        p = kb / name
        if isinstance(content, bytes):
            p.write_bytes(content)
        else:
            p.write_text(content, encoding="utf-8")
    if embed_fn == "unset":
        build_kb.build(kb)
    else:
        build_kb.build(kb, embed_fn=embed_fn)
    return kb


def _rows(kb):
    conn = sqlite3.connect(str(kb / "kb.db"))
    try:
        return conn.execute("SELECT doc, heading, content FROM chunks").fetchall()
    finally:
        conn.close()


def _meta(kb):
    conn = sqlite3.connect(str(kb / "kb.db"))
    try:
        return dict(conn.execute("SELECT key, value FROM meta"))
    finally:
        conn.close()


def _vecs(kb):
    conn = sqlite3.connect(str(kb / "kb.db"))
    try:
        blobs = [r[0] for r in conn.execute(
            "SELECT vec FROM chunks WHERE vec IS NOT NULL").fetchall()]
    finally:
        conn.close()
    out = []
    for b in blobs:
        a = array.array("f")
        a.frombytes(b)
        out.append(a)
    return out


_TOOL_SEQ = [0]


def _load_pack_tool(rel: Path):
    _TOOL_SEQ[0] += 1
    spec = importlib.util.spec_from_file_location(
        f"kbtool_{rel.stem}_{_TOOL_SEQ[0]}", rel)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _make_tool(kb, monkeypatch, which="common"):
    env_key = "KB_COMMON_DIR" if which == "common" else "KB_HR_DIR"
    monkeypatch.setenv(env_key, str(kb))
    return _load_pack_tool(PACK_TOOLS[which])


def _out_docs(out: str) -> list[str]:
    """从检索输出按出现顺序提取来源文档名。"""
    docs = []
    for line in out.splitlines():
        if line.startswith("【来源："):
            docs.append(line[len("【来源："):].split(" ·")[0])
    return docs


def _embed_from_map(mapping):
    """确定性假向量：按 chunk 内容查表；缺表项即测试构造错误。"""
    def _fn(texts):
        return [mapping[t] for t in texts]
    return _fn


def _block_network(monkeypatch):
    def _blocked(*args, **kwargs):
        raise OSError("network disabled by test")
    monkeypatch.setattr(socket, "socket", _blocked)
    monkeypatch.setattr(socket, "create_connection", _blocked)


# ------------------------------------------------------------ 样例文件生成


def _mk_docx(path: Path):
    from docx import Document
    d = Document()
    d.add_heading("年假制度", level=1)
    d.add_paragraph("连续工作满一年的员工，可享受五天带薪年假。")
    d.add_heading("考勤规定", level=2)
    d.add_paragraph("工作日九点前打卡。")
    d.save(path)


def _mk_pdf(path: Path, pages_text: list[str]):
    """手工构造最小 PDF（Helvetica 单行文本，仅 ASCII），供 pypdf 提取。"""
    def esc(s: str) -> str:
        return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    objs: dict[int, bytes] = {}
    n = len(pages_text)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objs[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objs[2] = f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode()
    objs[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    for i, txt in enumerate(pages_text):
        pn, cn = 4 + 2 * i, 5 + 2 * i
        objs[pn] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 3 0 R >> >> /Contents {cn} 0 R >>").encode()
        stream = f"BT /F1 12 Tf 72 720 Td ({esc(txt)}) Tj ET".encode()
        objs[cn] = (f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for num in sorted(objs):
        offsets[num] = len(out)
        out += f"{num} 0 obj\n".encode() + objs[num] + b"\nendobj\n"
    xref = len(out)
    maxnum = max(offsets)
    out += f"xref\n0 {maxnum + 1}\n".encode() + b"0000000000 65535 f \n"
    for i in range(1, maxnum + 1):
        out += f"{offsets[i]:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {maxnum + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF").encode()
    path.write_bytes(bytes(out))


def _mk_xlsx(path: Path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "假期额度"
    ws.append(["假别", "天数"])
    ws.append(["年假", "5"])
    ws.append(["病假", "10"])
    wb.save(path)


def _mk_pptx(path: Path):
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "入职流程"
    s.placeholders[1].text = "第一天报到并领取设备"
    prs.save(path)


# ---------------------------------------------------------------- AC1 五格式


def test_ac1_docx(tmp_path):
    kb = tmp_path / "kb"
    kb.mkdir()
    _mk_docx(kb / "制度.docx")
    build_kb.build(kb, embed_fn=None)
    rows = _rows(kb)
    headings = {h for _, h, _ in rows}
    assert "年假制度" in headings
    assert any("五天带薪年假" in c for _, h, c in rows if h == "年假制度")
    assert "考勤规定" in headings


def test_ac1_pdf(tmp_path):
    kb = tmp_path / "kb"
    kb.mkdir()
    _mk_pdf(kb / "handbook.pdf", ["Annual leave five paid days per year",
                                  "Check in before nine oclock"])
    build_kb.build(kb, embed_fn=None)
    rows = _rows(kb)
    headings = {h for _, h, _ in rows}
    assert headings == {"第1页", "第2页"}
    assert any("Annual leave" in c for _, _, c in rows)


def test_ac1_xlsx(tmp_path):
    kb = tmp_path / "kb"
    kb.mkdir()
    _mk_xlsx(kb / "额度.xlsx")
    build_kb.build(kb, embed_fn=None)
    rows = _rows(kb)
    assert {h for _, h, _ in rows} == {"假期额度"}
    assert any("年假" in c and "5" in c for _, _, c in rows)


def test_ac1_pptx(tmp_path):
    kb = tmp_path / "kb"
    kb.mkdir()
    _mk_pptx(kb / "培训.pptx")
    build_kb.build(kb, embed_fn=None)
    rows = _rows(kb)
    assert "入职流程" in {h for _, h, _ in rows}
    assert any("第一天报到" in c for _, _, c in rows)


def test_ac1_html(tmp_path):
    kb = tmp_path / "kb"
    kb.mkdir()
    (kb / "报销.html").write_text(
        "<html><body><h1>报销制度</h1><p>发票需在当月内提交</p></body></html>",
        encoding="utf-8")
    build_kb.build(kb, embed_fn=None)
    rows = _rows(kb)
    assert "报销制度" in {h for _, h, _ in rows}
    assert any("当月内提交" in c for _, _, c in rows)


# ---------------------------------------------------------------- AC2 不回退


def test_ac2_md_txt_unchanged(tmp_path, monkeypatch):
    kb = _build_kb(tmp_path, {
        "信息安全.md": "# 保密协议\n\n新员工入职一周内签署保密协议。",
        "纪要.txt": "会议纪要：下周三评审方案。",
    }, embed_fn=None)
    rows = _rows(kb)
    assert ("信息安全.md", "保密协议",
            "新员工入职一周内签署保密协议。") in rows
    assert any(h == "（前言）" and "下周三评审方案" in c for _, h, c in rows)
    assert build_kb.tokenize("年假") == ["年假"]  # 中文 bigram 分词不变

    tool = _make_tool(kb, monkeypatch)
    out = tool.search_common_knowledge("保密协议")
    assert out.startswith("【来源：信息安全.md · 保密协议】")


# ---------------------------------------------------------------- AC3 向量通道


def test_ac3_vecs_stored(tmp_path):
    kb = _build_kb(tmp_path, {
        "x.md": "# X\n\nleave leave",
        "y.md": "# Y\n\ndog cat",
    }, embed_fn=_embed_from_map({"leave leave": [1.0, 0.0], "dog cat": [0.0, 1.0]}))
    meta = _meta(kb)
    assert meta["has_vec"] == "1"
    assert meta["dim"] == "2"
    vecs = _vecs(kb)
    assert len(vecs) == len(_rows(kb))
    assert all(len(v) == 2 for v in vecs)


def test_ac3_query_vector_ranking(tmp_path, monkeypatch):
    """BM25 零命中时按点积排序返回（向量通道独立可用）。"""
    kb = _build_kb(tmp_path, {
        "x.md": "# X\n\nalpha one",
        "y.md": "# Y\n\nalpha two",
    }, embed_fn=_embed_from_map({"alpha one": [1.0, 0.0], "alpha two": [0.0, 1.0]}))
    tool = _make_tool(kb, monkeypatch)
    monkeypatch.setattr(tool, "_embed_texts", lambda texts: [[1.0, 0.0]])
    out = tool.search_common_knowledge("zulu")  # zulu 无词项命中 → 向量通道
    assert _out_docs(out) == ["x.md", "y.md"]


# ---------------------------------------------------------------- AC4 RRF 与退化


def _rrf_kb(tmp_path):
    """三块文档：BM25 名次 a>b>c（leave 词频 4/2/1），向量名次 b>c>a。"""
    return _build_kb(tmp_path, {
        "a.md": "# A\n\nleave leave leave leave",
        "b.md": "# B\n\nleave leave",
        "c.md": "# C\n\nleave",
    }, embed_fn=_embed_from_map({
        "leave leave leave leave": [1.0, 0.0],   # 点积 1.0 → 向量第 3
        "leave leave": [1.0, 1.0],               # 点积 2.0 → 向量第 1
        "leave": [1.0, 0.5],                     # 点积 1.5 → 向量第 2
    }))


def test_ac4_rrf_fusion(tmp_path, monkeypatch):
    kb = _rrf_kb(tmp_path)
    tool = _make_tool(kb, monkeypatch)
    monkeypatch.setattr(tool, "_embed_texts", lambda texts: [[1.0, 1.0]])
    # RRF(k=60)：b=1/62+1/61 最高；a=1/61+1/63 次之；c=1/63+1/62 最低
    out = tool.search_common_knowledge("leave")
    assert _out_docs(out) == ["b.md", "a.md", "c.md"]


def test_ac4_degrade_to_bm25(tmp_path, monkeypatch):
    """未启用向量（has_vec=0）→ 纯 BM25 排序照常工作。"""
    kb = _build_kb(tmp_path, {
        "a.md": "# A\n\nleave leave leave leave",
        "b.md": "# B\n\nleave leave",
        "c.md": "# C\n\nleave",
    }, embed_fn=None)
    assert _meta(kb)["has_vec"] == "0"
    tool = _make_tool(kb, monkeypatch)
    out = tool.search_common_knowledge("leave")
    assert _out_docs(out) == ["a.md", "b.md", "c.md"]


def test_ac4_degrade_to_vector(tmp_path, monkeypatch):
    """查询词无词项命中（BM25 通道空）→ 自动退化为向量排序。"""
    kb = _rrf_kb(tmp_path)
    tool = _make_tool(kb, monkeypatch)
    monkeypatch.setattr(tool, "_embed_texts", lambda texts: [[1.0, 1.0]])
    out = tool.search_common_knowledge("zulu")
    assert _out_docs(out) == ["b.md", "c.md", "a.md"]


def test_ac4_embed_fail_falls_back(tmp_path, monkeypatch):
    """查询侧向量引擎不可用（返回 None）→ 退化 BM25 不报错。"""
    kb = _rrf_kb(tmp_path)
    tool = _make_tool(kb, monkeypatch)
    monkeypatch.setattr(tool, "_embed_texts", lambda texts: None)
    out = tool.search_common_knowledge("leave")
    assert _out_docs(out) == ["a.md", "b.md", "c.md"]


# ---------------------------------------------------------------- AC6 拒绝路径


def test_ac6_legacy_format_rejected(tmp_path, capsys):
    kb = _build_kb(tmp_path, {
        "old.doc": b"\xd0\xcf\x11\xe0legacy-ole",
        "good.md": "# G\n\n正常内容一段。",
    }, embed_fn=None)
    out = capsys.readouterr().out
    assert "WARNING" in out
    assert "old.doc" in out and ".doc" in out and "转存" in out
    docs = {d for d, _, _ in _rows(kb)}
    assert docs == {"good.md"}                    # 其余文档不受阻


def test_ac6_corrupt_file_skipped(tmp_path, capsys):
    kb = _build_kb(tmp_path, {
        "broken.docx": b"not a zip archive",
        "good.md": "# G\n\n正常内容一段。",
    }, embed_fn=None)
    out = capsys.readouterr().out
    assert "WARNING" in out and "broken.docx" in out
    docs = {d for d, _, _ in _rows(kb)}
    assert docs == {"good.md"}


def test_ac6_scanned_pdf_skipped(tmp_path, capsys):
    kb = tmp_path / "kb"
    kb.mkdir()
    _mk_pdf(kb / "scan.pdf", [""])                # 无文本层 → 疑似扫描件
    (kb / "good.md").write_text("# G\n\n正常内容。", encoding="utf-8")
    build_kb.build(kb, embed_fn=None)
    out = capsys.readouterr().out
    assert "WARNING" in out and "scan.pdf" in out and "扫描件" in out
    docs = {d for d, _, _ in _rows(kb)}
    assert docs == {"good.md"}


def test_ac6_oversize_skipped(tmp_path, capsys, monkeypatch):
    monkeypatch.setattr(kb_parsers, "MAX_FILE_MB", 1)
    kb = _build_kb(tmp_path, {
        "big.txt": "x" * 1_200_000,
        "good.md": "# G\n\n正常内容。",
    }, embed_fn=None)
    out = capsys.readouterr().out
    assert "WARNING" in out and "big.txt" in out and "上限" in out
    docs = {d for d, _, _ in _rows(kb)}
    assert docs == {"good.md"}


# ---------------------------------------------------------------- AC7 零网络


def test_ac7_no_network(tmp_path, monkeypatch):
    """封禁 socket 后索引与查询全链路可用（注入假向量，不触发引擎加载）。"""
    _block_network(monkeypatch)
    kb = _build_kb(tmp_path, {
        "k.md": "# K\n\nleave go",
    }, embed_fn=_embed_from_map({"leave go": [1.0, 0.0]}))
    tool = _make_tool(kb, monkeypatch)
    monkeypatch.setattr(tool, "_embed_texts", lambda texts: [[1.0, 0.0]])
    out = tool.search_common_knowledge("leave")
    assert _out_docs(out) == ["k.md"]


# ---------------------------------------------------------------- AC8 自包含


def test_ac8_pack_tools_self_contained():
    """自包含静态断言：解析真实 import 语句（docstring 里的文字不误伤）。"""
    import ast
    for which, rel in PACK_TOOLS.items():
        tree = ast.parse(rel.read_text(encoding="utf-8"))
        banned = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                for a in node.names:
                    if a.name.split(".")[0] in ("server", "tools"):
                        banned.append(a.name)
            elif isinstance(node, ast.ImportFrom) and node.level == 0 and node.module:
                if node.module.split(".")[0] in ("server", "tools"):
                    banned.append(node.module)
        assert not banned, f"{which} 包工具不得 import server/tools：{banned}"


# ---------------------------------------------------------------- 引擎配置


def test_embed_openai_backend_build(tmp_path, monkeypatch):
    """EMBED_BACKEND=openai：外部 OpenAI 兼容 /v1/embeddings，按输入顺序返回向量。"""
    monkeypatch.setenv("EMBED_BACKEND", "openai")
    monkeypatch.setenv("EMBED_API_BASE", "http://emb.local/v1")
    monkeypatch.setenv("EMBED_API_KEY", "k-test")

    class _Resp:
        def read(self):
            return json.dumps({"data": [
                {"embedding": [0.5, 0.5]},   # 第 1 块（apple pie）
                {"embedding": [0.1, 0.9]},   # 第 2 块（dog cat）
            ]}).encode()

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

    import urllib.request
    monkeypatch.setattr(urllib.request, "urlopen", lambda req, timeout=None: _Resp())
    kb = _build_kb(tmp_path, {
        "a.md": "# A\n\napple pie",
        "b.md": "# B\n\ndog cat",
    })
    meta = _meta(kb)
    assert meta["has_vec"] == "1" and meta["dim"] == "2"
    assert meta["embed_backend"] == "openai"
    vecs = _vecs(kb)
    assert len(vecs) == 2
    assert abs(vecs[0][0] - 0.5) < 1e-6 and abs(vecs[1][1] - 0.9) < 1e-6


def test_embed_engine_missing_falls_back(tmp_path, capsys, monkeypatch):
    """local 引擎不可用（未装 fastembed/缺模型）→ WARNING 并退化为 BM25-only 索引。"""
    def _boom():
        raise ImportError("fastembed is not installed")
    monkeypatch.setattr(build_kb, "_get_engine", _boom)
    kb = _build_kb(tmp_path, {"a.md": "# A\n\napple pie"})   # 走 env 默认 local
    out = capsys.readouterr().out
    assert "WARNING" in out and "向量" in out
    assert _meta(kb)["has_vec"] == "0"
    assert len(_rows(kb)) == 1                                # BM25 索引照常


# ---------------------------------------------------------------- AC5 门控真模型


@pytest.mark.skipif(os.environ.get("RUN_EMBED_TESTS") != "true",
                    reason="门控 RUN_EMBED_TESTS：需 fastembed 与已缓存模型文件")
def test_ac5_semantic_query_offline(tmp_path, monkeypatch):
    """真模型（bge-small-zh）：语义问法命中 + 封网后全离线推理。"""
    pytest.importorskip("fastembed")
    kb = tmp_path / "kb"
    monkeypatch.setenv("KB_HR_DIR", str(kb))
    tool = _load_pack_tool(PACK_TOOLS["hr"])
    build_kb._get_engine()          # 预热两侧引擎（此时允许联网查缓存）
    tool._get_engine()
    _block_network(monkeypatch)     # 封网：索引与查询必须全离线
    kb.mkdir(parents=True)
    (kb / "假期管理规定.md").write_text(
        "# 假期管理规定\n\n员工连续工作满十二个月，可享受五天带薪年假。", encoding="utf-8")
    (kb / "考勤制度.md").write_text(
        "# 考勤制度\n\n工作日实行打卡考勤，迟到早退按规则处理。", encoding="utf-8")
    build_kb.build(kb)              # 默认 local 引擎（已缓存，离线推理）
    out = tool.search_hr_knowledge("新员工入职后什么时候可以休假")
    assert out.startswith("【来源：假期管理规定.md")


# ------------------------------------- SEC-7 原子重建（spec：security-hardening-v0.17.32 AC7）


def _docs(db_path):
    conn = sqlite3.connect(str(db_path))
    try:
        return {r[0] for r in conn.execute("SELECT DISTINCT doc FROM chunks")}
    finally:
        conn.close()


def test_ac7_build_preserves_previous_as_bak(tmp_path):
    """SEC-7：重建产出新库，旧库保留为 kb.db.bak（只留一版可回滚）；
    首次构建无上一版，不产生 bak。"""
    kb = tmp_path / "kb"
    kb.mkdir()
    (kb / "a.md").write_text("# A\n内容甲", encoding="utf-8")
    build_kb.build(kb, embed_fn=None)
    assert _docs(kb / "kb.db") == {"a.md"}
    assert not (kb / "kb.db.bak").exists()           # 首建无上一版
    (kb / "b.md").write_text("# B\n内容乙", encoding="utf-8")
    build_kb.build(kb, embed_fn=None)
    assert _docs(kb / "kb.db") == {"a.md", "b.md"}   # 新库含增量
    assert _docs(kb / "kb.db.bak") == {"a.md"}       # 上一版可回滚


def test_ac7_build_failure_keeps_old_db(tmp_path, monkeypatch):
    """SEC-7：构建中途异常 → 原库保持完好、无临时残留
    （旧实现「先 unlink 再重建」在失败时丢失旧索引——本用例为其回归防线）。"""
    kb = tmp_path / "kb"
    kb.mkdir()
    (kb / "a.md").write_text("# A\n内容甲", encoding="utf-8")
    build_kb.build(kb, embed_fn=None)
    (kb / "b.md").write_text("# B\n内容乙", encoding="utf-8")

    def boom(*a, **k):
        raise RuntimeError("注入失败：落库段异常")

    monkeypatch.setattr(build_kb.sqlite3, "connect", boom)
    with pytest.raises(RuntimeError):
        build_kb.build(kb, embed_fn=None)
    monkeypatch.undo()                               # 恢复真实 connect 以校验旧库
    assert _docs(kb / "kb.db") == {"a.md"}           # 旧库完好
    assert not (kb / "kb.db.tmp").exists()           # 无临时残留


def test_review_c4_build_exception_closes_conn_and_keeps_original(
        tmp_path, monkeypatch):
    """评审修复 C4 回归：connect 成功后落库段异常 → 连接被 finally 关闭
    （句柄不锁 tmp），原始异常正常传播（Windows 下不被 PermissionError
    掩盖）、旧库完好、无残留。（旧实现仅正常路径关闭连接，异常路径句柄
    未释放 → 外层 finally unlink 抛 PermissionError 掩盖原始异常，且 tmp
    残留 + CREATE TABLE 撞表会毒化下次构建。）"""
    kb = tmp_path / "kb"
    kb.mkdir()
    (kb / "a.md").write_text("# A\n内容甲", encoding="utf-8")
    build_kb.build(kb, embed_fn=None)
    (kb / "b.md").write_text("# B\n内容乙", encoding="utf-8")

    real_connect = build_kb.sqlite3.connect

    class _BoomConn:
        def __init__(self, path):
            self._c = real_connect(path)
            self.closed = False

        def execute(self, *a, **k):
            return self._c.execute(*a, **k)

        def executemany(self, *a, **k):
            raise RuntimeError("注入失败：落库段异常")

        def commit(self, *a, **k):
            self._c.commit(*a, **k)

        def close(self):
            self.closed = True
            self._c.close()

    conn = {}

    def _fake_connect(path, *a, **k):
        conn["h"] = _BoomConn(path)
        return conn["h"]

    monkeypatch.setattr(build_kb.sqlite3, "connect", _fake_connect)
    with pytest.raises(RuntimeError, match="落库段异常"):
        build_kb.build(kb, embed_fn=None)
    monkeypatch.undo()
    assert conn["h"].closed                        # 异常路径同样关闭连接
    assert _docs(kb / "kb.db") == {"a.md"}          # 旧库完好
    assert not (kb / "kb.db.tmp").exists()          # 无残留（句柄已释放）


def test_ac6_replace_failure_keeps_db_in_place(tmp_path, monkeypatch):
    """v0.17.33 AC6：新库替换现役库 3 连败 → 现役库从未被移走（备份由
    copy 生成），旧版仍为现役、bak 保留上一版、无 tmp 残留——不存在
    『bak 已转存、现役缺失』状态，不再需要回滚分支。"""
    kb = tmp_path / "kb"
    kb.mkdir()
    (kb / "a.md").write_text("# A\n内容甲", encoding="utf-8")
    build_kb.build(kb, embed_fn=None)
    (kb / "b.md").write_text("# B\n内容乙", encoding="utf-8")

    real_replace = os.replace

    def _flaky_replace(src, dst):
        if str(src).endswith("kb.db.tmp") and str(dst).endswith("kb.db"):
            raise OSError("注入失败：现役替换被锁")
        return real_replace(src, dst)

    monkeypatch.setattr(build_kb.os, "replace", _flaky_replace)
    with pytest.raises(OSError, match="现役替换被锁"):
        build_kb.build(kb, embed_fn=None)
    monkeypatch.undo()
    assert _docs(kb / "kb.db") == {"a.md"}          # 现役库在原位（从未移走）
    assert _docs(kb / "kb.db.bak") == {"a.md"}      # 备份已由 copy 生成
    assert not (kb / "kb.db.tmp").exists()          # 无临时残留


def test_ac6_interrupt_window_db_always_present(tmp_path, monkeypatch):
    """v0.17.33 AC6：备份后、现役替换前注入 BaseException（模拟 kill，
    任何 except/回滚分支都不执行）→ 现役库必须仍在原位——复验实测旧
    实现此处 kb.db 消失只剩 bak（两次 os.replace 之间的中断窗口）。"""
    kb = tmp_path / "kb"
    kb.mkdir()
    (kb / "a.md").write_text("# A\n内容甲", encoding="utf-8")
    build_kb.build(kb, embed_fn=None)
    (kb / "b.md").write_text("# B\n内容乙", encoding="utf-8")

    real_replace = os.replace

    def _kill_at_replace(src, dst):
        if str(src).endswith("kb.db.tmp") and str(dst).endswith("kb.db"):
            raise KeyboardInterrupt("模拟进程中断")
        return real_replace(src, dst)

    monkeypatch.setattr(build_kb.os, "replace", _kill_at_replace)
    with pytest.raises(KeyboardInterrupt):
        build_kb.build(kb, embed_fn=None)
    assert _docs(kb / "kb.db") == {"a.md"}          # 现役库完好
    assert _docs(kb / "kb.db.bak") == {"a.md"}      # 备份在位
